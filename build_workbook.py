"""Build Product_Count_Analysis.xlsx using COUNTIF / COUNTIFS formulas (live formulas, not pasted values)."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC = "data/superstore_sample.csv"
OUT = "excel/Product_Count_Analysis.xlsx"

NAVY, ORANGE, LIGHT = "0B1F4B", "F97316", "EAF0FA"
hdr_font = Font(bold=True, color="FFFFFF", name="Calibri")
hdr_fill = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="CBD5E1")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, box
        cell.alignment = Alignment(horizontal="center", vertical="center")

with open(SRC, newline="") as f:
    data = list(csv.reader(f))
header, rows = data[0], data[1:]
num_cols = {"Row ID", "Sales", "Quantity", "Discount", "Profit"}

wb = Workbook()

# ---------- Sheet 1: Data ----------
ws = wb.active
ws.title = "Data"
ws.append(header)
for r in rows:
    out = []
    for h, v in zip(header, r):
        if h in ("Row ID", "Quantity"): out.append(int(v))
        elif h in ("Sales", "Discount", "Profit"): out.append(float(v))
        else: out.append(v)
    ws.append(out)
last = ws.max_row
style_header(ws, 1, 1, len(header))
for i, h in enumerate(header, 1):
    ws.column_dimensions[get_column_letter(i)].width = max(12, min(28, len(h) + 6))
ws.column_dimensions["O"].width = 30
ws.freeze_panes = "A2"
tab = Table(displayName="Orders", ref=f"A1:{get_column_letter(len(header))}{last}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
ws.add_table(tab)

col = {h: get_column_letter(i) for i, h in enumerate(header, 1)}
CAT = f"Data!${col['Category']}$2:${col['Category']}${last}"
SUB = f"Data!${col['Sub-Category']}$2:${col['Sub-Category']}${last}"
REG = f"Data!${col['Region']}$2:${col['Region']}${last}"

# ---------- Sheet 2: Unique Products ----------
up = wb.create_sheet("Unique Products")
seen = {}
ci, si, pi, ni = (header.index(x) for x in ("Category", "Sub-Category", "Product ID", "Product Name"))
for r in rows:
    seen.setdefault(r[pi], (r[ci], r[si], r[ni]))
up.append(["Product ID", "Category", "Sub-Category", "Product Name"])
for pid, (c, s, n) in sorted(seen.items(), key=lambda kv: (kv[1][0], kv[1][1], kv[0])):
    up.append([pid, c, s, n])
style_header(up, 1, 1, 4)
for i, w in enumerate([16, 18, 18, 34], 1):
    up.column_dimensions[get_column_letter(i)].width = w
up.freeze_panes = "A2"
ulast = up.max_row
UCAT = f"'Unique Products'!$B$2:$B${ulast}"
USUB = f"'Unique Products'!$C$2:$C${ulast}"

categories = ["Office Supplies", "Furniture", "Technology"]
subcats = {}
for r in rows:
    subcats.setdefault(r[ci], set()).add(r[si])

# ---------- Sheet 3: Product Count ----------
pc = wb.create_sheet("Product Count")
pc["A1"] = "Product Count by Category"
pc["A1"].font = Font(bold=True, size=16, color=NAVY)
pc["A2"] = "All figures are live COUNTIF / COUNTIFS formulas on the Data and Unique Products sheets."
pc["A2"].font = Font(italic=True, color="64748B")

pc.append([])
pc.append(["Category", "Order Lines (COUNTIF)", "Unique Products (COUNTIF)", "% of Order Lines", "Rank"])
style_header(pc, 4, 1, 5)
for i, c in enumerate(categories):
    r = 5 + i
    pc[f"A{r}"] = c
    pc[f"B{r}"] = f'=COUNTIF({CAT},A{r})'
    pc[f"C{r}"] = f"=COUNTIF({UCAT},A{r})"
    pc[f"D{r}"] = f"=B{r}/$B$8"
    pc[f"E{r}"] = f"=RANK(B{r},$B$5:$B$7)"
    pc[f"D{r}"].number_format = "0.0%"
pc["A8"], pc["B8"], pc["C8"], pc["D8"] = "Total", "=SUM(B5:B7)", "=SUM(C5:C7)", "=SUM(D5:D7)"
pc["D8"].number_format = "0.0%"
for c in "ABCDE":
    pc[f"{c}8"].font = Font(bold=True)
    pc[f"{c}8"].fill = PatternFill("solid", fgColor=LIGHT)
for r in range(5, 9):
    for c in "ABCDE":
        pc[f"{c}{r}"].border = box
pc["A9"] = "Check: total order lines vs. data rows"
pc["B9"] = f"=IF(B8=COUNTA(Data!$A$2:$A${last}),\"OK\",\"MISMATCH\")"
pc["A9"].font = Font(italic=True, color="64748B")

# Sub-category table (COUNTIFS)
pc["A11"] = "Sub-Category Breakdown (COUNTIFS)"
pc["A11"].font = Font(bold=True, size=13, color=NAVY)
pc.append([])
pc["A12"], pc["B12"], pc["C12"], pc["D12"] = "Category", "Sub-Category", "Order Lines", "Unique Products"
style_header(pc, 12, 1, 4)
r = 13
for c in categories:
    for s in sorted(subcats[c]):
        pc[f"A{r}"], pc[f"B{r}"] = c, s
        pc[f"C{r}"] = f'=COUNTIFS({CAT},A{r},{SUB},B{r})'
        pc[f"D{r}"] = f"=COUNTIFS({UCAT},A{r},{USUB},B{r})"
        for cc in "ABCD":
            pc[f"{cc}{r}"].border = box
        r += 1
sub_end = r - 1

# Category x Region matrix
mr = sub_end + 3
pc[f"A{mr-1}"] = "Order Lines by Category and Region (COUNTIFS)"
pc[f"A{mr-1}"].font = Font(bold=True, size=13, color=NAVY)
regions = ["West", "East", "Central", "South"]
pc[f"A{mr}"] = "Category"
for j, rg in enumerate(regions):
    pc.cell(row=mr, column=2 + j, value=rg)
pc.cell(row=mr, column=6, value="Total")
style_header(pc, mr, 1, 6)
for i, c in enumerate(categories):
    rr = mr + 1 + i
    pc[f"A{rr}"] = c
    for j in range(4):
        L = get_column_letter(2 + j)
        pc[f"{L}{rr}"] = f'=COUNTIFS({CAT},$A{rr},{REG},{L}${mr})'
    pc[f"F{rr}"] = f"=SUM(B{rr}:E{rr})"
    for cc in "ABCDEF":
        pc[f"{cc}{rr}"].border = box
for i, w in enumerate([22, 24, 26, 18, 10, 10], 1):
    pc.column_dimensions[get_column_letter(i)].width = w

ch = BarChart()
ch.type, ch.title = "col", "Order Lines by Category"
ch.add_data(Reference(pc, min_col=2, min_row=4, max_row=7), titles_from_data=True)
ch.set_categories(Reference(pc, min_col=1, min_row=5, max_row=7))
ch.y_axis.title, ch.legend = "Count", None
ch.height, ch.width = 7.5, 13
pc.add_chart(ch, "H4")

# ---------- Sheet 4: Top Category ----------
tc = wb.create_sheet("Top Category")
tc["A1"] = "Largest Product Category"
tc["A1"].font = Font(bold=True, size=16, color=NAVY)
items = [
    ("Top category (by order lines)", "=INDEX('Product Count'!$A$5:$A$7,MATCH(MAX('Product Count'!$B$5:$B$7),'Product Count'!$B$5:$B$7,0))"),
    ("Order lines in top category", "=MAX('Product Count'!$B$5:$B$7)"),
    ("Share of all order lines", "=B5/'Product Count'!$B$8"),
    ("Top category (by unique products)", "=INDEX('Product Count'!$A$5:$A$7,MATCH(MAX('Product Count'!$C$5:$C$7),'Product Count'!$C$5:$C$7,0))"),
    ("Unique products in that category", "=MAX('Product Count'!$C$5:$C$7)"),
    ("Largest sub-category (order lines)", "=INDEX('Product Count'!$B$13:$B$%d,MATCH(MAX('Product Count'!$C$13:$C$%d),'Product Count'!$C$13:$C$%d,0))" % (sub_end, sub_end, sub_end)),
    ("Order lines in that sub-category", "=MAX('Product Count'!$C$13:$C$%d)" % sub_end),
]
tc["A3"], tc["B3"] = "Metric", "Value"
style_header(tc, 3, 1, 2)
for i, (k, f) in enumerate(items):
    tc[f"A{4+i}"], tc[f"B{4+i}"] = k, f
    tc[f"A{4+i}"].border = tc[f"B{4+i}"].border = box
    tc[f"B{4+i}"].alignment = Alignment(horizontal="left")
tc["B6"].number_format = "0.0%"
tc["B4"].font = Font(bold=True, color=ORANGE, size=12)
tc.column_dimensions["A"].width = 38
tc.column_dimensions["B"].width = 30

# ---------- Sheet 5: Notes ----------
nt = wb.create_sheet("Notes")
notes = [
    "Task 15 - Product Count Analysis (Data Analytics track)",
    "",
    "COUNTIF  : counts rows that meet ONE condition, e.g. =COUNTIF(Category, \"Furniture\").",
    "COUNTIFS : counts rows that meet MANY conditions at once, e.g. Category AND Region.",
    "Order lines = number of rows (a product bought in 10 orders = 10 lines).",
    "Unique products = each Product ID counted once (list on 'Unique Products' sheet).",
    "Data/superstore_sample.csv is Superstore-style sample data. To use the real dataset, paste it into the 'Data' sheet",
    "(same column names) - every formula range will need to be extended to the new last row.",
]
for i, t in enumerate(notes, 1):
    nt[f"A{i}"] = t
nt["A1"].font = Font(bold=True, size=14, color=NAVY)
nt.column_dimensions["A"].width = 110

wb.move_sheet("Product Count", offset=-2)
wb.move_sheet("Top Category", offset=-2)
wb.active = 0
wb.save(OUT)
print("saved", OUT, "| sub_end", sub_end)
